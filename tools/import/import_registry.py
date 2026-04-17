from openpyxl import load_workbook
from pathlib import Path
import csv
import uuid

BASE_DIR = Path(__file__).resolve().parent
EXCEL_FILE = BASE_DIR / "Пост-мониторинг свод на 14.04.xlsx"
OUTPUT_CSV = BASE_DIR / "registry_import_staging.csv"

BATCH_ID = str(uuid.uuid4())
print("BATCH_ID:", BATCH_ID)

wb = load_workbook(EXCEL_FILE, read_only=True, data_only=True)
ws = wb["перечень"]  # основной лист

def norm(v):
    if v is None:
        return ""
    return str(v).strip()

header_row_idx = 2
rows_iter = ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx, values_only=True)
header_excel = [norm(v) for v in next(rows_iter)]

header_map = {
    "П/п": "row_number",
    "Анализ / мониторингB2:N155M146B2:O16B2:N161": "record_type",
    "ЦИКЛ": "cycle",
    "Сфера": "sphere",
    "Предложения": "proposal_text",
    "Ответственный исполнитель": "responsible_org",
    "Заинтересованные государственные органы": "interested_orgs",
    "Форма завершения": "completion_form",
    "Срок исполнения": "due_raw",
    "Статус ГО": "status_raw",
    "Позиция ГО на 2024-2025гг.": "position_go_2024_2025",
    "Позиция ГО на 27.03.2026": "position_go_2026_03_27",
    "Позиция АДГС": "position_adgs",
    "Кейс": "case_raw",
}

columns = ["batch_id"] + list(header_map.values())

with OUTPUT_CSV.open("w", encoding="utf-8", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=columns)
    writer.writeheader()

    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        cells = [norm(v) for v in row]
        if not any(cells):
            continue

        rec = {"batch_id": BATCH_ID}
        for src, dst in header_map.items():
            try:
                idx = header_excel.index(src)
            except ValueError:
                continue
            rec[dst] = cells[idx] if idx < len(cells) else ""

        if rec.get("row_number"):
            try:
                rec["row_number"] = int(float(rec["row_number"]))
            except ValueError:
                pass

        writer.writerow(rec)

print("Wrote CSV:", OUTPUT_CSV)
